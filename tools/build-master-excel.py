"""Builds the layout of "AB Capital Master Data" (sheets, headers, dropdowns, editable cells),
prefilled with the dashboard's sample data. tools/build-master-excel.ps1 then adds the
Generate SQL macro and button, protects it and saves it as .xlsm.

Usage: python tools/build-master-excel.py <sample.json> <out.xlsx>
"""
import json
import sys
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

EDIT_ROWS = 2000  # editable rows per sheet

TEAL = "1D7A72"
INK = "14302D"
EDIT_FILL = PatternFill("solid", fgColor="FFF8DC")   # editable cells
HEAD_FILL = PatternFill("solid", fgColor=TEAL)
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="E2E1DC")
UNLOCKED = Protection(locked=False)

LISTS = {
    "Departments": ["Management", "Sales & Marketing", "Company Formation", "PRO & Visa", "Banking",
                    "Accounting & Tax", "Loans & Mortgage", "Advisory", "HR & Admin"],
    "AccessRoles": ["admin", "manager", "dept_head", "staff", "finance", "viewer"],
    "Platforms": ["Meta", "Google", "LinkedIn", "TikTok", "Other"],
    "LeadSources": ["Website", "WhatsApp", "Referral", "Book-a-Call", "Cost Calculator", "Walk-in", "Meta Ads", "Google Ads", "Other"],
    "LeadStatus": ["New", "Contacted", "Proposal", "Won", "Lost"],
    "CaseStatus": ["Open", "On Hold", "Completed", "Cancelled"],
    "BankStatus": ["Preparing", "Submitted", "Under Review", "Approved", "Rejected", "Withdrawn"],
    "TaxTypes": ["VAT Return", "Corporate Tax Return", "E-invoicing", "Tax Residency Certificate", "Other"],
    "ExpiryTypes": ["Trade Licence", "Investor Visa", "Employment Visa", "Golden Visa", "Emirates ID", "Establishment Card"],
}

# (header, width, kind, list name or None, note)
# kind: text | date | month | money | int
SHEETS = {
    "Services": [
        ("Department", 22, "text", "Departments", "Department that delivers the service"),
        ("Service", 30, "text", None, "Service name, unique. Used by Leads and Cases"),
        ("DefaultFeeAED", 16, "money", None, None),
    ],
    "Team": [
        ("Staff", 20, "text", None, "Full name, unique. Used as Owner / Assignee"),
        ("Department", 22, "text", "Departments", None),
        ("JobTitle", 26, "text", None, None),
        ("AccessRole", 14, "text", "AccessRoles", "Dashboard access level (default staff)"),
        ("LeaveFrom", 14, "date", None, "Leave start, if on leave"),
        ("LeaveTo", 14, "date", None, "Leave end"),
    ],
    "Clients": [
        ("ClientID", 12, "text", None, "Unique ID such as C001. Used on every other sheet"),
        ("Name", 22, "text", None, "Main contact person"),
        ("Company", 30, "text", None, None),
        ("Jurisdiction", 20, "text", None, "Dubai Mainland, SRTIP, IFZA ..."),
        ("LicenceNo", 14, "text", None, None),
        ("TRN", 20, "text", None, "Tax registration number"),
        ("StartDate", 14, "date", None, None),
    ],
    "Campaigns": [
        ("Month", 12, "month", None, "Any date in the month"),
        ("Campaign", 32, "text", None, "Campaign name, same spelling as on Leads"),
        ("Platform", 12, "text", "Platforms", None),
        ("SpendAED", 14, "money", None, "Amount spent that month"),
        ("Impressions", 14, "int", None, None),
        ("Clicks", 12, "int", None, None),
    ],
    "Leads": [
        ("LeadID", 10, "text", None, "Unique, such as L0001"),
        ("Date", 13, "date", None, None),
        ("Source", 16, "text", "LeadSources", None),
        ("Campaign", 32, "text", None, "For Meta Ads / Google Ads leads: name from the Campaigns sheet"),
        ("ClientName", 22, "text", None, None),
        ("Service", 28, "text", None, "Name from the Services sheet"),
        ("ValueAED", 13, "money", None, None),
        ("Status", 12, "text", "LeadStatus", None),
        ("Owner", 16, "text", None, "Name from the Team sheet"),
    ],
    "Cases": [
        ("CaseID", 10, "text", None, "Unique, such as K0001"),
        ("ClientID", 10, "text", None, "ID from the Clients sheet"),
        ("Service", 28, "text", None, "Name from the Services sheet"),
        ("Stage", 18, "text", None, "Current step, e.g. Documents, Stamping"),
        ("Assignee", 16, "text", None, "Name from the Team sheet"),
        ("OpenedOn", 13, "date", None, None),
        ("DueDate", 13, "date", None, None),
        ("ClosedOn", 13, "date", None, "Leave blank while open"),
        ("FeeAED", 12, "money", None, None),
        ("Status", 12, "text", "CaseStatus", None),
    ],
    "Expiries": [
        ("ClientID", 10, "text", None, "ID from the Clients sheet"),
        ("Type", 20, "text", "ExpiryTypes", None),
        ("Number", 14, "text", None, None),
        ("ExpiryDate", 13, "date", None, None),
    ],
    "Banking": [
        ("ClientID", 10, "text", None, "ID from the Clients sheet"),
        ("Bank", 16, "text", None, None),
        ("AppliedOn", 13, "date", None, None),
        ("Status", 14, "text", "BankStatus", None),
        ("RejectionReason", 28, "text", None, None),
    ],
    "TaxFilings": [
        ("ClientID", 10, "text", None, "ID from the Clients sheet"),
        ("Type", 24, "text", "TaxTypes", None),
        ("Period", 12, "text", None, "Q3 2026, FY 2025 ..."),
        ("DueDate", 13, "date", None, None),
        ("FiledOn", 13, "date", None, "Leave blank until filed"),
    ],
    "Finance": [
        ("Month", 12, "month", None, "Any date in the month (from Tally)"),
        ("RevenueAED", 14, "money", None, None),
        ("ExpensesAED", 14, "money", None, None),
        ("CollectedAED", 14, "money", None, None),
        ("VATPayableAED", 15, "money", None, None),
    ],
    "Receivables": [
        ("ClientName", 30, "text", None, "As written in Tally"),
        ("InvoiceNo", 16, "text", None, "Unique"),
        ("InvoiceDate", 13, "date", None, None),
        ("DueDate", 13, "date", None, None),
        ("AmountAED", 14, "money", None, "Unpaid amount"),
    ],
    "BankBalances": [
        ("Bank", 16, "text", None, None),
        ("Account", 22, "text", None, None),
        ("BalanceAED", 14, "money", None, None),
        ("AsOf", 13, "date", None, None),
    ],
}

FORMATS = {"date": "dd-mmm-yyyy", "month": "mmm-yyyy", "money": "#,##0", "int": "#,##0", "text": "@"}


def to_date(v):
    if not v:
        return None
    return datetime.strptime(v[:10], "%Y-%m-%d").date()


def rows_from_sample(s):
    today = date.today()
    staff_leave = (today - timedelta(days=1), today + timedelta(days=5))
    return {
        "Services": [[x["Department"], x["Service"], x["Fee"]] for x in s["_services"]],
        "Team": [[m["Staff"], m["Department"], m["Role"], "manager" if m["Department"] == "Management" else "staff",
                  staff_leave[0] if m["OnLeave"] == "Y" else None, staff_leave[1] if m["OnLeave"] == "Y" else None] for m in s["Team"]],
        "Clients": [[c["ClientID"], c["Name"], c["Company"], c["Jurisdiction"], c["LicenceNo"], c["TRN"] or None, to_date(c["StartDate"])] for c in s["Clients"]],
        "Campaigns": [[to_date(c["Month"]), c["Campaign"], c["Platform"], c["SpendAED"], c["Impressions"], c["Clicks"]] for c in s["Campaigns"]],
        "Leads": [[l["LeadID"], to_date(l["Date"]), l["Source"], l["Campaign"] or None, l["ClientName"], l["Service"], l["ValueAED"], l["Status"], l["Owner"]] for l in s["Leads"]],
        "Cases": [[c["CaseID"], c["ClientID"], c["Service"], c["Stage"], c["Assignee"], to_date(c["OpenedOn"]), to_date(c["DueDate"]),
                   to_date(c["ClosedOn"]), c["FeeAED"], c["Status"]] for c in s["Cases"]],
        "Expiries": [[e["ClientID"], e["Type"], e["Number"], to_date(e["ExpiryDate"])] for e in s["Expiries"]],
        "Banking": [[b["ClientID"], b["Bank"], to_date(b["AppliedOn"]), b["Status"], b["RejectionReason"] or None] for b in s["Banking"]],
        "TaxFilings": [[t["ClientID"], t["Type"], t["Period"], to_date(t["DueDate"]),
                        (to_date(t["DueDate"]) - timedelta(days=3)) if t["Filed"] == "Y" else None] for t in s["TaxCalendar"]],
        "Finance": [[to_date(f["Month"]), f["RevenueAED"], f["ExpensesAED"], f["CollectedAED"], f["VATPayableAED"]] for f in s["Finance"]],
        "Receivables": [[r["ClientName"], r["InvoiceNo"], to_date(r["InvoiceDate"]), to_date(r["DueDate"]), r["AmountAED"]] for r in s["Receivables"]],
        "BankBalances": [[b["Bank"], b["Account"], b["BalanceAED"], to_date(b["AsOf"])] for b in s["BankBalances"]],
    }


def build_start(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 90
    ws["B2"] = "AB Capital - Master Data"
    ws["B2"].font = Font(size=20, bold=True, color=INK)
    ws["B3"] = "Personalize Business Solutions  |  Source for the CEO dashboard"
    ws["B3"].font = Font(size=10, bold=True, color=TEAL)
    steps = [
        ("1. Edit", "Change values only in the light-yellow cells on each sheet. Headers and everything else are locked."),
        ("2. Rules", "Use the dropdowns. Dates as dates, amounts in AED. IDs and names must match across sheets (ClientID, Service, Staff, Campaign)."),
        ("3. Add / remove", "Add a row by typing in the next empty yellow row. Remove a row by clearing its cells."),
        ("4. Generate", "Click Generate SQL below. It checks the data and writes db\\sample-data.sql in this folder."),
        ("5. Publish", "In a terminal in the project folder: git add db/sample-data.sql, then git commit -m \"Update dashboard data\", then git push. The site updates in about a minute."),
        ("Keep in place", "Keep this workbook in the project folder (next to index.html). Open it with macros enabled."),
        ("Warning", "db/sample-data.sql is public on GitHub Pages. Keep real client data out of it while the repo is public."),
    ]
    for i, (k, v) in enumerate(steps):
        r = 5 + i
        ws.cell(r, 2, k).font = Font(bold=True, color=INK)
        ws.cell(r, 3, v).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30
    ws["B16"] = "Last generated"
    ws["B16"].font = Font(bold=True, color=INK)
    ws["C16"].number_format = "dd-mmm-yyyy hh:mm"
    ws["C16"].alignment = Alignment(horizontal="left")


def build_lists(wb):
    ws = wb.create_sheet("Lists")
    refs = {}
    for c, (name, values) in enumerate(LISTS.items(), start=1):
        ws.cell(1, c, name).font = Font(bold=True)
        for r, v in enumerate(values, start=2):
            ws.cell(r, c, v)
        col = get_column_letter(c)
        refs[name] = f"Lists!${col}$2:${col}${len(values) + 1}"
    return refs


def build_sheet(wb, name, cols, rows, refs):
    ws = wb.create_sheet(name)
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(cols))
    for c, (head, width, kind, lst, note) in enumerate(cols, start=1):
        cell = ws.cell(1, c, head)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center")
        if note:
            cell.comment = Comment(note, "AB Capital")
        col = get_column_letter(c)
        ws.column_dimensions[col].width = width
        rng = f"{col}2:{col}{EDIT_ROWS + 1}"
        if lst:
            dv = DataValidation(type="list", formula1=f"={refs[lst]}", allow_blank=True)
            if lst == "ExpiryTypes":
                dv.errorStyle = "warning"  # suggestions only; other types allowed
            dv.error, dv.errorTitle = "Pick a value from the list.", head
        elif kind in ("date", "month"):
            dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
            dv.error, dv.errorTitle = "Enter a date, e.g. 25-Oct-2026.", head
        elif kind in ("money", "int"):
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
            dv.error, dv.errorTitle = "Enter a number (0 or more).", head
        else:
            dv = None
        if dv:
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add(rng)
    ws.row_dimensions[1].height = 22
    ws.auto_filter.ref = f"A1:{last_col}{EDIT_ROWS + 1}"

    for r in range(2, EDIT_ROWS + 2):
        data = rows[r - 2] if r - 2 < len(rows) else None
        for c, (_, _, kind, _, _) in enumerate(cols, start=1):
            cell = ws.cell(r, c)
            if data is not None and data[c - 1] not in (None, ""):
                v = data[c - 1]
                cell.value = str(v) if kind == "text" else v
            cell.protection = UNLOCKED
            cell.fill = EDIT_FILL
            cell.number_format = FORMATS[kind]
            cell.border = Border(bottom=THIN)


def main(src, out):
    with open(src, encoding="utf-8") as f:
        sample = json.load(f)
    data = rows_from_sample(sample)
    wb = Workbook()
    build_start(wb.active)
    wb.active.title = "Start"
    refs = build_lists(wb)
    for name, cols in SHEETS.items():
        build_sheet(wb, name, cols, data[name], refs)
    # Lists last in the tab order; the macro step hides it
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    wb.save(out)
    print(f"{out} written")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
